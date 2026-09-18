"""Build the frozen LiuQP/ellipsoid/VCC mathematical specification as DOCX.

The source of truth is the adjacent Markdown file.  This converter intentionally
keeps LaTeX equations as centered, editable Cambria Math text so every symbol is
preserved without relying on a network equation renderer.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "UR5_incremental_depth_LiuQP_ellipsoid_VCC_math_spec.md"
OUTPUT = ROOT / "UR5增量深度点云_LiuQP椭球VCC数学规格与实验协议.docx"


def set_run_font(run, ascii_name: str, east_asia_name: str, size: float | None = None):
    run.font.name = ascii_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia_name)
    if size is not None:
        run.font.size = Pt(size)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


INLINE = re.compile(r"(\*\*.*?\*\*|`.*?`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, *, default_size: float = 10.5) -> None:
    pos = 0
    for match in INLINE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, "Times New Roman", "宋体", default_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            set_run_font(run, "Times New Roman", "宋体", default_size)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.color.rgb = RGBColor(80, 80, 80)
            set_run_font(run, "Consolas", "等线", default_size - 0.5)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(f"{label}（{url}）")
            run.font.color.rgb = RGBColor(31, 78, 121)
            run.underline = True
            set_run_font(run, "Times New Roman", "宋体", default_size)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, "Times New Roman", "宋体", default_size)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for name, size, color, before, after in (
        ("Title", 22, "17365D", 0, 18),
        ("Heading 1", 16, "17365D", 18, 8),
        ("Heading 2", 13, "1F4E79", 14, 6),
        ("Heading 3", 11.5, "365F91", 10, 4),
        ("Heading 4", 10.5, "44546A", 8, 3),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)

    for name in ("List Bullet", "List Number"):
        styles[name].font.name = "Times New Roman"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        styles[name].font.size = Pt(10.5)

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Cm(0.7)
    code.paragraph_format.right_indent = Cm(0.4)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(3)
    code.paragraph_format.first_line_indent = Cm(0)

    equation = styles.add_style("Equation Text", WD_STYLE_TYPE.PARAGRAPH)
    equation.font.name = "Cambria Math"
    equation._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    equation.font.size = Pt(9.5)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(4)
    equation.paragraph_format.space_after = Pt(4)
    equation.paragraph_format.keep_together = True
    equation.paragraph_format.first_line_indent = Cm(0)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, "Arial", "微软雅黑", 9)


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在 Word 打开时更新"
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2, placeholder, fld_char3])


def add_table(doc: Document, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(width):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            text = row[c_idx] if c_idx < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, text, default_size=8.5)
            if r_idx == 0:
                shade_cell(cell, "D9EAF7")
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_markdown(doc: Document, lines: list[str]) -> None:
    i = 0
    in_code = False
    equation: list[str] | None = None
    while i < len(lines):
        raw = lines[i].rstrip("\n")
        stripped = raw.strip()

        if equation is not None:
            if stripped == "\\]":
                p = doc.add_paragraph(style="Equation Text")
                run = p.add_run(" ".join(equation))
                set_run_font(run, "Cambria Math", "Cambria Math", 9.5)
                equation = None
            else:
                equation.append(stripped)
            i += 1
            continue
        if stripped == "\\[":
            equation = []
            i += 1
            continue
        if stripped.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph(style="Code Block")
            p.add_run(raw)
            i += 1
            continue
        if not stripped or stripped == "---":
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines):
            divider = lines[i + 1].strip()
            if divider.startswith("|") and re.fullmatch(r"[|:\- ]+", divider):
                rows = []
                rows.append([x.strip() for x in stripped.strip("|").split("|")])
                i += 2
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rows.append([x.strip() for x in lines[i].strip().strip("|").split("|")])
                    i += 1
                add_table(doc, rows)
                continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, text, default_size=22)
                p.paragraph_format.space_before = Pt(70)
                p.paragraph_format.space_after = Pt(24)
            else:
                p = doc.add_paragraph(style=f"Heading {level - 1}")
                add_inline(p, text, default_size={2: 16, 3: 13, 4: 11.5}[level])
            i += 1
            continue

        if stripped.startswith("- [ ]"):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, "☐ " + stripped[5:].strip())
            i += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, stripped[2:].strip())
            i += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            # Preserve the source number literally. Word's built-in list style
            # otherwise continues numbering across unrelated sections.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.first_line_indent = Cm(-0.74)
            add_inline(p, f"{numbered.group(1)}. {numbered.group(2)}")
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1


def main() -> None:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("UR5 增量深度点云 · LiuQP 椭球/VCC 技术规格")
    run.font.color.rgb = RGBColor(100, 100, 100)
    set_run_font(run, "Arial", "微软雅黑", 8.5)
    add_page_number(section.footer.paragraphs[0])

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    parse_markdown(doc, lines)

    # Put a generated-on notice at the end, not in the research claims.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("Word 版由同名 Markdown 单一来源自动生成；公式以可编辑 LaTeX 文本保留。")
    run.italic = True
    run.font.color.rgb = RGBColor(100, 100, 100)
    set_run_font(run, "Times New Roman", "宋体", 8.5)

    # Add a TOC section after the title page using a new section at document end
    # would disturb source order; instead prepend a compact TOC field after the
    # first metadata paragraphs through XML insertion.
    title_p = doc.paragraphs[0]
    toc_heading = OxmlElement("w:p")
    toc_p = doc.add_paragraph()
    toc_p.style = doc.styles["Heading 1"]
    toc_p.add_run("目录")
    add_toc(doc.add_paragraph())
    # Move the two TOC paragraphs immediately after the first five source blocks.
    body = doc._element.body
    toc_nodes = [doc.paragraphs[-2]._p, doc.paragraphs[-1]._p]
    insert_after = doc.paragraphs[4]._p if len(doc.paragraphs) > 4 else title_p._p
    for node in reversed(toc_nodes):
        insert_after.addnext(node)

    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
